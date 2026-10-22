"""Validate an Excel workbook (.xlsx) against an `excel` data contract (Stage 2).

The workbook is read with openpyxl in read-only mode, one sheet (`format.sheet`, else the
first), with the header on `format.header_row` (default 1; rows above it, such as a title,
are ignored). Cells become text the way a CSV export would write them, and then go
through exactly the CSV engine's header, row and file checks (`csvfeed.check_rows`), so a
problem gets the same reason code whichever format it arrived in:

  empty cell              ""
  whole number            "12" (Excel stores every number as a float; 12.0 becomes "12")
  other number            the shortest repr, "12.5"
  date or midnight time   "2026-10-22"; any other datetime "2026-10-22T08:30:00"
  boolean                 "true" / "false"
  formula                 its cached value. A formula never calculated (a file written by
                          a library, not saved by Excel) has none and reads as empty, so
                          a `required` column catches it

Reason codes added by this format:
  FILE.INVALID_XLSX       not a readable .xlsx workbook (a bad zip, broken XML, anything
                          openpyxl cannot parse): quarantined, never a crashed task
  FILE.MACROS             the workbook contains a VBA project (.xlsm content): refused,
                          whatever the file is called
  FILE.TOO_LARGE          over MAX_UNCOMPRESSED bytes unzipped, or a sheet that spans more
                          than MAX_CELLS cells. Checked before any row is read: a tiny file
                          with one cell at XFD1048576 would otherwise expand to billions of
                          empty cells
  SCHEMA.MISSING_SHEET:<name>   no such sheet, or it is a chart, not a worksheet
  VALUE.EXCEL_ERROR:<col> (row level) a formula error such as #DIV/0! or #N/A as the
                          cell's value
"""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, date, datetime, time
from typing import Any

from dip.csvfeed import CsvResult, check_rows
from dip.fhir import Violation

MAX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_CELLS = 5_000_000
EXCEL_ERRORS = frozenset(
    {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA",
     "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#UNKNOWN!", "#CONNECT!", "#BUSY!"}
)  # fmt: skip


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def _rows(data: bytes, fmt: dict[str, Any], result: CsvResult) -> list[list[str]] | None:
    import openpyxl

    def refuse(code: str, detail: str) -> None:
        result.file_violations.append(Violation(code, detail))

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
        if any(i.filename.lower().endswith("vbaproject.bin") for i in infos):
            refuse("FILE.MACROS", "the workbook contains a VBA project (macros)")
            return None
        unzipped = sum(i.file_size for i in infos)
        if unzipped > MAX_UNCOMPRESSED:
            refuse("FILE.TOO_LARGE", f"{unzipped} bytes unzipped, limit {MAX_UNCOMPRESSED}")
            return None
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # any parse failure is a bad file, never a crashed task
        refuse("FILE.INVALID_XLSX", f"not an .xlsx workbook: {type(exc).__name__}: {exc}")
        return None
    try:
        name = fmt.get("sheet")
        names = [ws.title for ws in wb.worksheets]  # worksheets only, not chartsheets
        if name is not None and name not in names:
            refuse(f"SCHEMA.MISSING_SHEET:{name}", f"worksheets are {names}")
            return None
        if not names:
            refuse("SCHEMA.MISSING_SHEET:", "the workbook has no worksheet")
            return None
        ws = wb[name] if name is not None else wb.worksheets[0]
        rows, cols = ws.max_row or 0, ws.max_column or 0
        if rows * cols > MAX_CELLS:
            refuse("FILE.TOO_LARGE", f"the sheet spans {rows} x {cols} cells, limit {MAX_CELLS}")
            return None
        raw, cells = [], 0
        for row in ws.iter_rows(values_only=True):
            cells += len(row)
            if cells > MAX_CELLS:  # a sheet without (or with a false) dimension record
                refuse("FILE.TOO_LARGE", f"more than {MAX_CELLS} cells")
                return None
            raw.append([cell_text(v) for v in row])
    except Exception as exc:
        refuse("FILE.INVALID_XLSX", f"unreadable sheet: {type(exc).__name__}: {exc}")
        return None
    finally:
        wb.close()
    return raw[fmt.get("header_row", 1) - 1 :]


def _error_cells(rows: list[list[str]]) -> dict[int, Violation]:
    """Data rows (1-based, as check_rows numbers them) holding a formula error."""
    if not rows:
        return {}
    header = rows[0]
    out = {}
    for n, row in enumerate(rows[1:], start=1):
        for col, value in zip(header, row, strict=False):
            if value in EXCEL_ERRORS:
                out[n] = Violation(f"VALUE.EXCEL_ERROR:{col}", f"{col} is the error {value}")
                break
    return out


def _normalize(rows: list[list[str]], result: CsvResult) -> list[list[str]]:
    """Excel pads every row to the sheet's widest column. Trailing empty header cells are
    not columns; data rows are cut back to the header's width when the extra cells are
    empty (a value out there stays, and fails ROW.FIELD_COUNT) and padded when short.
    Rows with no value at all are blank lines: dropped and counted, as in CSV."""
    if not rows:
        return rows
    header = list(rows[0])
    while header and header[-1] == "":
        header.pop()
    out = [header]
    blank = 0
    for row in rows[1:]:
        if not any(row):
            blank += 1
            continue
        row = list(row)
        while len(row) > len(header) and row[-1] == "":
            row.pop()
        row += [""] * (len(header) - len(row))
        out.append(row)
    result.skipped["blank_line"] = blank
    return out


def validate_excel(data: bytes, contract: Any, received_at: datetime | None = None) -> CsvResult:
    spec = contract.spec
    today = (received_at or datetime.now(UTC)).astimezone(UTC).date()
    result = CsvResult()
    rows = _rows(data, spec["format"], result)
    if rows is None:
        return result
    rows = _normalize(rows, result)
    return check_rows(rows, spec, today, result, row_violations=_error_cells(rows))
